"""Build the GymBuddy Lab 1 presentation (editable .pptx) -> docs/GymBuddy_Lab1_Presentation.pptx"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
MOCK = os.path.join(HERE, "..", "design", "mockups")

NAVY = RGBColor(0x11, 0x18, 0x27)
ORANGE = RGBColor(0xF9, 0x73, 0x16)
ORANGE_DK = RGBColor(0xC2, 0x41, 0x0C)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TEAM = [
    ("Aida Ismayilova", "57689", "Product Owner + Developer",
     "Problem statement, MVP, user stories and priorities; owns the board content."),
    ("Ozodbek Tursunaliyev", "57264", "Tech Lead + Developer",
     "Data model, architecture, Android and Firebase setup, auth feature."),
    ("Elturan Valizada", "58990", "QA / DevOps + Developer",
     "Repository, branch ruleset, PR reviews, board hygiene, testing."),
]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def run(p, text, size, color=INK, bold=False):
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return r


def box(slide, l, t, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, l, t, w, h)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def tf_of(shape, margin=0.15):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.08)
    return tf


def slide(title, subtitle=None):
    s = prs.slides.add_slide(BLANK)
    bar = box(s, 0, 0, SW, Inches(1.2), NAVY, shape=MSO_SHAPE.RECTANGLE)
    tf = tf_of(bar, 0.5)
    run(tf.paragraphs[0], title, 32, WHITE, True)
    if subtitle:
        run(tf.add_paragraph(), subtitle, 15, RGBColor(0xFD, 0xBA, 0x74))
    box(s, 0, Inches(1.2), SW, Inches(0.07), ORANGE, shape=MSO_SHAPE.RECTANGLE)
    return s


def bullets(s, items, l, t, w, h, size=20):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run(p, "•  ", size, ORANGE, True)
        if isinstance(it, tuple):
            run(p, it[0], size, INK, True)
            run(p, " " + it[1], size, MUTED)
        else:
            run(p, it, size, INK)


def table(s, rows, l, t, w, col_w, size=15, row_h=0.5):
    shp = s.shapes.add_table(len(rows), len(rows[0]), l, t, w, Inches(row_h * len(rows)))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (LIGHT if i % 2 else WHITE)
            run(c.text_frame.paragraphs[0], val, size, WHITE if i == 0 else INK, i == 0)
    return tbl


# 1 Title -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, SW, SH, NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, Inches(0.8), Inches(1.5), Inches(1.6), Inches(1.6), ORANGE)
t = box(s, Inches(0.8), Inches(1.5), Inches(1.6), Inches(1.6))
tf_of(t).paragraphs[0].alignment = PP_ALIGN.CENTER
run(t.text_frame.paragraphs[0], "GB", 54, WHITE, True)
tb = s.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(11.5), Inches(3.5))
tf = tb.text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "GymBuddy", 66, WHITE, True)
run(tf.add_paragraph(), "Log your workouts. Find your training partner.", 26, RGBColor(0xFD, 0xBA, 0x74))
p = tf.add_paragraph(); p.space_before = Pt(24)
run(p, "Team Project – Mobile Application in the Cloud · Lab 1", 18, RGBColor(0xCB, 0xD5, 0xE1))
p = tf.add_paragraph()
run(p, "Aida Ismayilova (57689)  ·  Ozodbek Tursunaliyev (57264)  ·  Elturan Valizada (58990)", 18, WHITE)

# 2 What is GymBuddy ------------------------------------------------------
s = slide("What is GymBuddy?", "An Android app backed by the cloud")
bullets(s, [
    ("Gym + Buddy:", "your workout log and your training partner finder in one app."),
    ("Log:", "record each workout (type, duration, notes) and see your history and weekly progress."),
    ("Find:", "post or join open training sessions, so you never have to train alone."),
    ("Cloud:", "everything is stored in Firebase, so your data follows you to any phone."),
], Inches(0.7), Inches(1.7), Inches(7.2), Inches(5), 22)
b = box(s, Inches(8.3), Inches(1.9), Inches(4.4), Inches(4.6), LIGHT)
tf = tf_of(b, 0.3)
run(tf.paragraphs[0], "Why the name?", 22, ORANGE_DK, True)
p = tf.add_paragraph(); p.space_before = Pt(10)
run(p, "“Buddy” is the gym word for a training partner. People who train with a partner skip fewer sessions – "
       "so the name says what the app is for: staying consistent, together.", 18, INK)

# 2b Where it can be used -------------------------------------------------
s = slide("Where GymBuddy can be used", "The same app helps different groups")
table(s, [
    ["Who", "How GymBuddy helps"],
    ["Individual gym-goers", "Keep a simple training history and see weekly progress"],
    ["Friends and training groups", "Plan shared sessions and see who joined"],
    ["Newcomers to a gym", "Find a partner instead of training alone"],
    ["University sports clubs", "Organise open sessions for students and staff"],
    ["Workplace wellness programmes", "Let colleagues arrange group workouts"],
    ["Small gyms and fitness studios", "Give members a free way to meet training partners"],
], Inches(0.7), Inches(1.8), Inches(11.9), [4.4, 7.5], 18, 0.62)
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.8))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "People who track their training see progress, and people who train with a partner skip fewer sessions.", 16, MUTED)

# 3 Problem ---------------------------------------------------------------
s = slide("The problem", "One sentence: who has what problem, and why it matters")
b = box(s, Inches(0.7), Inches(1.8), Inches(11.9), Inches(1.5), ORANGE)
tf = tf_of(b, 0.4)
run(tf.paragraphs[0], "“Gym-goers forget what they trained and struggle to find training partners, "
                      "so progress stalls and motivation drops.”", 24, WHITE, True)
bullets(s, [
    ("Today:", "workouts live in notes apps, memory or nowhere; partners are found by chance."),
    ("Impact:", "no visible progress and no accountability lead to people quitting."),
    ("Our answer:", "a simple cloud log plus a place to arrange sessions with others."),
], Inches(0.7), Inches(3.7), Inches(11.9), Inches(3), 22)

# 4 MVP -------------------------------------------------------------------
s = slide("MVP features", "Sprint 1 = Labs 1–2, the rest is planned for later")
table(s, [
    ["#", "Feature", "Sprint"],
    ["1", "Register, log in, log out, auto-login (Firebase Authentication)", "Sprint 1 (Lab 1)"],
    ["2", "List my workouts", "Sprint 1 (Lab 2)"],
    ["3", "Add, edit and delete a workout (detail/edit screen)", "Sprint 1 (Lab 2)"],
    ["4", "Weekly stats calculated by a Cloud Function", "Later (Lab 3)"],
    ["5", "Find a partner: create and join training sessions", "Later"],
], Inches(0.7), Inches(1.8), Inches(11.9), [0.7, 8.2, 3.0], 18, 0.62)
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.0), Inches(11.9), Inches(1))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Limits we respect: 5 screens (incl. login and register) and max. 3 Firestore collections.", 16, MUTED)

# 5 Screens ---------------------------------------------------------------
s = slide("How the app will look", "5 screens: Login · Register · Workouts · Workout detail/edit · Partners")
s.shapes.add_picture(os.path.join(MOCK, "00_overview.png"), Inches(0.5), Inches(1.7), width=Inches(12.3))
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.55), Inches(11.9), Inches(0.6))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Concept mockups. Login, Register and a temporary Home are built in Lab 1; the rest follow in Labs 2–3.", 15, MUTED)

# 6 Data model ------------------------------------------------------------
s = slide("Data model", "Cloud Firestore · 3 collections · ownerId links every document to its owner")
table(s, [
    ["Collection", "Fields (type)"],
    ["users", "uid (doc id) · email (string) · displayName (string) · createdAt (timestamp)"],
    ["workouts", "ownerId (string) · title (string) · type (string) · durationMin (number) · "
                 "notes (string) · date (timestamp) · createdAt (timestamp)"],
    ["sessions", "ownerId (string) · title (string) · gym (string) · startsAt (timestamp) · "
                 "maxPeople (number) · participantIds (array of string)"],
], Inches(0.7), Inches(1.8), Inches(11.9), [2.2, 9.7], 17, 0.95)
tb = s.shapes.add_textbox(Inches(0.7), Inches(5.9), Inches(11.9), Inches(1))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Security Rules (Lab 2) make sure users can only edit their own data – enforced in Firebase, not just in the app.", 17, ORANGE_DK, True)

# 7 Architecture ----------------------------------------------------------
s = slide("Architecture", "Screens never talk to Firebase directly")
layers = [("UI", "LoginScreen · RegisterScreen · HomeScreen · AppNavHost"),
          ("Logic", "AuthValidator · AuthViewModel (StateFlow)"),
          ("Data", "AuthRepository → FirebaseAuthRepository")]
y = 1.7
for name, desc in layers:
    b = box(s, Inches(0.7), Inches(y), Inches(6.6), Inches(1.15), LIGHT, ORANGE)
    tf = tf_of(b, 0.25)
    run(tf.paragraphs[0], name, 20, ORANGE_DK, True)
    run(tf.add_paragraph(), desc, 15, INK)
    y += 1.45
for i in range(2):
    a = box(s, Inches(3.75), Inches(2.85 + i * 1.45), Inches(0.5), Inches(0.3), ORANGE, shape=MSO_SHAPE.DOWN_ARROW)
cloud = [("Android app", "Kotlin + Jetpack Compose", NAVY),
         ("Firebase Authentication", "email + password", ORANGE),
         ("Cloud Firestore", "workouts, sessions, users", ORANGE),
         ("Cloud Functions", "Lab 3 – weekly stats", MUTED)]
y = 1.7
for name, sub, col in cloud:
    b = box(s, Inches(8.3), Inches(y), Inches(4.3), Inches(0.95), col)
    tf = tf_of(b, 0.25)
    run(tf.paragraphs[0], name, 18, WHITE, True)
    run(tf.add_paragraph(), sub, 13, WHITE)
    y += 1.1
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.2), Inches(11.9), Inches(1))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Why: simple screens, and the logic (e.g. AuthValidator) can be unit-tested on the JVM in Lab 3.", 16, MUTED)

# 8 User stories ----------------------------------------------------------
s = slide("User stories", "Written as GitHub Issues, in the format required by the lab instructions")
b = box(s, Inches(0.7), Inches(1.7), Inches(5.9), Inches(4.9), LIGHT)
tf = tf_of(b, 0.3)
run(tf.paragraphs[0], "Title: As a gym-goer, I want to log in with email and password", 15, INK, True)
for txt, bold in [("", False), ("User story", True),
                  ("As a gym-goer, I want to log in with my email and password, so that I can access my own workouts.", False),
                  ("", False), ("Acceptance criteria", True),
                  ("☐ Valid credentials open the home screen", False),
                  ("☐ Wrong credentials show an error message", False),
                  ("☐ A logged-in user skips the login screen after restart", False)]:
    run(tf.add_paragraph(), txt, 14, INK, bold)
bullets(s, [
    ("Format:", "title + user story + acceptance criteria checklist."),
    ("Labels:", "user-story plus android / firebase / ui."),
    ("Board:", "Sprint 1 issues in the Sprint 1 column, others in Backlog; each has an assignee."),
    ("Closing:", "the PR says “Closes #issue”, so the issue moves to Done on merge."),
], Inches(6.9), Inches(1.7), Inches(5.9), Inches(5), 17)

# 9 Sprint 1 stories ------------------------------------------------------
s = slide("Sprint 1 backlog", "Each story has acceptance criteria and an assignee")
table(s, [
    ["User story", "Labels", "Assignee"],
    ["As a gym-goer, I want to register and log in with email and password", "user-story · android · firebase", "Ozodbek"],
    ["As a gym-goer, I want to see a list of my workouts", "user-story · android · ui", "Elturan"],
    ["As a gym-goer, I want to add a workout", "user-story · android · firebase", "Aida"],
    ["As a gym-goer, I want to edit or delete a workout", "user-story · android · firebase", "Elturan"],
], Inches(0.7), Inches(1.8), Inches(11.9), [6.6, 3.9, 1.4], 16, 0.8)
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.0), Inches(11.9), Inches(1))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Backlog: weekly stats Cloud Function, find a partner (sessions).", 17, MUTED)

# 10 Team ----------------------------------------------------------------
s = slide("Team and roles", "Everyone writes code – roles say who is accountable for a decision")
x = 0.7
for name, sid, role, duty in TEAM:
    b = box(s, Inches(x), Inches(1.8), Inches(3.8), Inches(4.3), LIGHT, ORANGE)
    tf = tf_of(b, 0.25)
    run(tf.paragraphs[0], name, 22, INK, True)
    run(tf.add_paragraph(), f"Student ID {sid}", 14, MUTED)
    p = tf.add_paragraph(); p.space_before = Pt(14)
    run(p, role, 18, ORANGE_DK, True)
    p = tf.add_paragraph(); p.space_before = Pt(10)
    run(p, duty, 16, INK)
    x += 4.05
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.3), Inches(11.9), Inches(0.8))
tb.text_frame.word_wrap = True
run(tb.text_frame.paragraphs[0], "Three people cover four roles: all three are Developers; the commit history will show everyone.", 16, MUTED)

# 11 Workflow -------------------------------------------------------------
s = slide("Our GitHub workflow", "How we work together")
steps = ["Issue on the board", "Feature branch", "Small commits", "Pull request", "Review + approve", "Squash merge → Done"]
x = 0.5
for i, st in enumerate(steps):
    b = box(s, Inches(x), Inches(2.0), Inches(1.95), Inches(1.3), ORANGE if i % 2 == 0 else NAVY)
    tf = tf_of(b, 0.1)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    run(tf.paragraphs[0], st, 15, WHITE, True)
    x += 2.1
bullets(s, [
    ("Protected main:", "ruleset protect-main – pull request required, no force push, no deletion."),
    ("Board columns:", "Backlog · Sprint 1 · In progress · In review · Done."),
    ("Templates:", "pull request template with “Closes #issue” and a test checklist."),
    ("Goal:", "at least two merged PRs, each reviewed by a teammate."),
], Inches(0.7), Inches(3.8), Inches(11.9), Inches(3), 19)

# 12 Demo -----------------------------------------------------------------
s = slide("Demo plan (3 minutes)", "Who shows what")
table(s, [
    ["Who", "Shows"],
    ["Aida (Product Owner)", "Problem statement and MVP features"],
    ["Ozodbek (Tech Lead)", "Data model and architecture diagram, then the live demo: register → logout → login → restart with auto-login"],
    ["Elturan (QA / DevOps)", "Board, merged PRs with review comments, branch ruleset"],
], Inches(0.7), Inches(1.8), Inches(11.9), [3.4, 8.5], 18, 0.85)
bullets(s, ["Expected question: “Where do you check that a user only edits their own data?” → "
            "In Firebase Security Rules (Lab 2), not only in the app."],
        Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.5), 18)

# 12b Status --------------------------------------------------------------
s = slide("Where we are now", "Lab 1 checklist: done and still to do")
table(s, [
    ["Done", "Still to do"],
    ["Problem statement, MVP, data model and architecture in the README",
     "Sprint board with 5 columns and issues in the right columns"],
    ["Public repo, first push to main, PR template",
     "Ruleset protecting main (PR required, no force push)"],
    ["Labels and 6 issues with user story and acceptance criteria",
     "Assignees on every Sprint 1 issue (teammates must accept their invitations)"],
    ["Authentication code: repository, validator, ViewModel, screens, navigation",
     "Firebase project, google-services.json, Email/Password sign-in, test on the emulator"],
    ["PR #7 merged (project overview)",
     "At least two merged PRs, each with a teammate review; Retro – Lab 1 issue"],
], Inches(0.7), Inches(1.8), Inches(11.9), [5.95, 5.95], 15, 0.85)

# 13 Next ----------------------------------------------------------------
s = slide("What's next", "Lab 1 delivers the walking skeleton")
bullets(s, [
    ("Lab 1:", "repo, board, Firebase project, register / login / logout / auto-login."),
    ("Lab 2:", "workouts in Firestore – list, add, edit, delete – and Security Rules."),
    ("Lab 3:", "Cloud Function for weekly stats, unit tests for the validator."),
    ("Retro – Lab 1:", "one issue, every member comments: went well / did not go well / we will change."),
], Inches(0.7), Inches(1.8), Inches(11.9), Inches(4), 22)
tb = s.shapes.add_textbox(Inches(0.7), Inches(6.0), Inches(11.9), Inches(1))
run(tb.text_frame.paragraphs[0], "Thank you – questions?", 30, ORANGE_DK, True)

for sl in prs.slides:
    for shp in sl.shapes:
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                if para.alignment is None:
                    para.alignment = PP_ALIGN.LEFT

out = os.path.join(HERE, "GymBuddy_Lab1_Presentation.pptx")
prs.save(out)
print("saved", out, len(prs.slides), "slides")

